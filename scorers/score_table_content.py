#!/usr/bin/env python3
"""
score_table_content.py
----------------------
Cell-content quality scorer for extracted tables.

Compares per-table hand-transcribed ground truth against extracted table
JSONs. GT sheets are named pNNN_II (page NNN, table index II) and map 1:1 to
table_pNNN_II.json. Run twice (old Camelot output, new Docling output) against
the same GT for the before/after comparison.

Scoring: normalized multiset match of non-empty cell strings per table.
  precision = matched / predicted_nonempty
  recall    = matched / gt_nonempty
  F1        = harmonic mean
Reports per-table and aggregate.

Usage:
  python3 score_table_content.py --gt hp_cell_gt.xlsx --tables output/tables --label docling
"""
import os, re, json, glob, argparse
from collections import Counter


def norm(s):
    s = str(s or "")
    s = re.sub(r"\.{3,}", " ", s)
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s


def cells_from_grid(ws):
    out = []
    for row in ws.iter_rows(values_only=True):
        for v in row:
            n = norm(v)
            # skip the trailing note row marker
            if n.startswith("^ correct"):
                continue
            if n:
                out.append(n)
    return Counter(out)


def load_gt(xlsx_path):
    """Return {(page, idx): Counter} keyed by table, from sheets named pNNN_II."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    gt = {}
    for ws in wb.worksheets:
        m = re.match(r"p(\d+)_(\d+)", ws.title.strip())
        if not m:
            continue
        key = (int(m.group(1)), int(m.group(2)))
        c = cells_from_grid(ws)
        if c:
            gt[key] = c
    return gt


def load_one_table(path):
    try:
        d = json.load(open(path, encoding="utf-8"))
    except Exception:
        return Counter()
    cells = []
    for h in d.get("headers", []):
        n = norm(h)
        if n and not n.isdigit():
            cells.append(n)
    for r in d.get("rows", []):
        vals = r.values() if isinstance(r, dict) else r
        for v in vals:
            n = norm(v)
            if n:
                cells.append(n)
    return Counter(cells)


def find_table_file(tables_dir, page, idx):
    """Exact match table_p{page}_{idx}.json; fall back to any table on page."""
    exact = os.path.join(tables_dir, f"table_p{page}_{idx}.json")
    if os.path.exists(exact):
        return exact, "exact"
    # fallback: a different index on the same page (old code may index differently)
    cand = sorted(glob.glob(os.path.join(tables_dir, f"table_p{page}_*.json")))
    if cand:
        return cand[0], "page-fallback"
    return None, "missing"


def prf(pred, gt):
    matched = sum((pred & gt).values())
    p_tot, g_tot = sum(pred.values()), sum(gt.values())
    P = matched / p_tot if p_tot else 0.0
    R = matched / g_tot if g_tot else 0.0
    F = 2*P*R/(P+R) if (P+R) else 0.0
    return P, R, F, matched, p_tot, g_tot


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--gt", required=True)
    ap.add_argument("--tables", required=True)
    ap.add_argument("--label", default="")
    args = ap.parse_args()

    gt = load_gt(args.gt)
    print(f"\n=== Cell-content scoring {('['+args.label+']') if args.label else ''} ===")
    print(f"GT tables: {len(gt)} | tables dir: {args.tables}\n")
    print(f"{'Table':>10} {'P':>6} {'R':>6} {'F1':>6}  {'match':>5}/{'pred':>5}/{'gt':>4}  match-type")

    tm = tp = tg = 0
    for (page, idx) in sorted(gt):
        path, how = find_table_file(args.tables, page, idx)
        pred = load_one_table(path) if path else Counter()
        P, R, F, m, pt, gtt = prf(pred, gt[(page, idx)])
        tm += m; tp += pt; tg += gtt
        print(f"{'p'+str(page)+'_'+str(idx):>10} {P:>6.3f} {R:>6.3f} {F:>6.3f}  {m:>5}/{pt:>5}/{gtt:>4}  {how}")

    AP = tm/tp if tp else 0.0
    AR = tm/tg if tg else 0.0
    AF = 2*AP*AR/(AP+AR) if (AP+AR) else 0.0
    print(f"\nAGGREGATE  P={AP:.3f}  R={AR:.3f}  F1={AF:.3f}  (matched {tm} / pred {tp} / gt {tg})\n")
    print("Note: 'page-fallback' means the exact table index wasn't found and a\n"
          "different table on that page was used — review those rows manually,\n"
          "especially for the OLD output where indices differ.\n")


if __name__ == "__main__":
    main()
